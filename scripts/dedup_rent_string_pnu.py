# -*- coding: utf-8 -*-
"""
원본 임대료 엑셀을 읽을 때 PNU코드를 문자열로 고정한 뒤,
중복키(PNU코드 + 매물게시일 + 계약면적 + 보증금 + 월세가) 기준으로 중복을 제거합니다.

기본 출력(권장, 빠름):
  - 2_processed/화성시_동탄구_상가데이터_임대료_dedup.csv   UTF-8 BOM
  - 2_processed/화성시_동탄구_상가데이터_임대료_dedup.parquet (pyarrow 필요)

선택:
  --xlsx : 동일 데이터를 xlsx로 저장 (26만 행 규모에서는 매우 느림)

후속 파이프라인에서는 CSV 또는 Parquet만 사용하는 것을 권장합니다.
"""

from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "1_raw" / "화성시_동탄구_상가데이터_임대료.xlsx"
OUT_DIR = ROOT / "2_processed"
OUT_CSV = OUT_DIR / "화성시_동탄구_상가데이터_임대료_dedup.csv"
OUT_PARQUET = OUT_DIR / "화성시_동탄구_상가데이터_임대료_dedup.parquet"
OUT_XLSX = OUT_DIR / "화성시_동탄구_상가데이터_임대료_dedup.xlsx"

DUP_KEYS = ["PNU코드", "매물게시일", "계약면적", "보증금", "월세가"]


def format_pnu_column_openpyxl(path: Path) -> None:
    """PNU코드 열에 텍스트 서식(@) 적용."""
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
    except ImportError as e:
        raise SystemExit("openpyxl이 필요합니다. pip install openpyxl") from e

    wb = load_workbook(path)
    ws = wb.active
    header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    col_idx = header.index("PNU코드") + 1
    ws.column_dimensions[get_column_letter(col_idx)].number_format = "@"
    wb.save(path)
    wb.close()


def dedupe_frame(df: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, int]]:
    df = df.copy()
    df["PNU코드"] = df["PNU코드"].astype(str).str.strip()
    df.loc[df["PNU코드"].isin(["nan", "None", ""]), "PNU코드"] = pd.NA

    n_total = len(df)
    df_ok = df.dropna(subset=DUP_KEYS)
    n_ok = len(df_ok)
    dropped_na = n_total - n_ok

    df_dedup = df_ok.drop_duplicates(subset=DUP_KEYS, keep="first")
    n_final = len(df_dedup)
    n_dup_removed = n_ok - n_final

    stats = {
        "n_total": n_total,
        "dropped_na": dropped_na,
        "n_ok": n_ok,
        "n_dup_removed": n_dup_removed,
        "n_final": n_final,
    }
    return df_dedup, stats


def save_outputs(df: pd.DataFrame, write_xlsx: bool) -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    # 1) CSV — 교환·열람용
    df.to_csv(OUT_CSV, index=False, encoding="utf-8-sig")

    # 2) Parquet — 분석·속도용 (PNU 문자열 유지)
    try:
        df.to_parquet(OUT_PARQUET, index=False, engine="pyarrow")
    except ImportError as e:
        raise SystemExit(
            "Parquet 저장에는 pyarrow가 필요합니다: pip install pyarrow\n"
            f"(CSV는 이미 저장됨: {OUT_CSV})"
        ) from e

    # 3) xlsx — 선택
    if write_xlsx:
        df.to_excel(OUT_XLSX, index=False, engine="openpyxl")
        format_pnu_column_openpyxl(OUT_XLSX)


def main() -> None:
    parser = argparse.ArgumentParser(description="임대료 원본 dedup (문자열 PNU)")
    parser.add_argument(
        "--xlsx",
        action="store_true",
        help="느린 xlsx도 함께 저장 (기본은 CSV + Parquet만)",
    )
    args = parser.parse_args()

    if not SRC.exists():
        raise SystemExit(f"원본 파일이 없습니다: {SRC}")

    df = pd.read_excel(SRC, dtype={"PNU코드": str})
    df_dedup, st = dedupe_frame(df)

    save_outputs(df_dedup, write_xlsx=args.xlsx)

    print(f"원본 행수: {st['n_total']:,}")
    print(f"중복키 결측 행(제외): {st['dropped_na']:,}")
    print(f"중복 제거 전(결측 제외): {st['n_ok']:,}")
    print(f"중복으로 제거된 행: {st['n_dup_removed']:,}")
    print(f"최종 행수: {st['n_final']:,}")
    if st["n_ok"]:
        print(f"제거율(중복): {100 * st['n_dup_removed'] / st['n_ok']:.2f}%")
    print(f"저장 CSV: {OUT_CSV}")
    print(f"저장 Parquet: {OUT_PARQUET}")
    if args.xlsx:
        print(f"저장 XLSX: {OUT_XLSX}")


if __name__ == "__main__":
    main()
